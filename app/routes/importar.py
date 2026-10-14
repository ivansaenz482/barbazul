"""
Importador de datos (migracion desde otro sistema).

Lee un archivo CSV o Excel (xlsx) y carga clientes o productos al sistema,
detectando duplicados (por cedula/RUC en clientes; por SKU o codigo de barras
en productos). Pensado para migrar la informacion de un sistema anterior.
"""
import csv
import io
import json
import os
import re
import unicodedata
import uuid

from flask import Blueprint, render_template, request, redirect, url_for, flash, current_app
from flask_login import login_required

from app.extensions import db
from app.models import Customer, Product
from app.utils import role_required
from app.audit import registrar

importar_bp = Blueprint("importar", __name__, url_prefix="/importar")

# Definicion de campos por tipo de importacion:
# (campo_destino, etiqueta, obligatorio, palabras_clave_para_detectar_columna)
CAMPOS = {
    "clientes": [
        ("full_name", "Nombre completo", True, ["nombre", "cliente", "name", "razon", "nombres", "apellidos", "razon_social"]),
        ("cedula", "Cédula / RUC", False, ["cedula", "ruc", "identificacion", "documento", "dni", "ci"]),
        ("business_name", "Nombre del negocio", False, ["negocio", "empresa", "business", "comercial"]),
        ("phone", "Teléfono / Celular", False, ["telefono", "celular", "phone", "movil", "whatsapp", "telf"]),
        ("email", "Correo", False, ["correo", "email", "mail"]),
        ("address", "Dirección", False, ["direccion", "address", "domicilio", "ubicacion"]),
    ],
    "productos": [
        ("name", "Nombre del producto", True, ["nombre", "producto", "descripcion", "name", "detalle", "articulo"]),
        ("sku", "Código / SKU", False, ["codigo", "sku", "cod", "barcode", "codigo_barras", "codbar", "referencia"]),
        ("cost_price", "Costo", False, ["costo", "cost", "precio_compra", "p_compra", "compra"]),
        ("sale_price", "Precio de venta", False, ["precio", "pvp", "precio_venta", "venta", "p_venta", "precio1"]),
        ("current_stock", "Stock / Existencia", False, ["stock", "cantidad", "existencia", "existencias", "saldo", "unidades"]),
    ],
}


def _norm(texto):
    s = (texto or "").strip().lower()
    s = "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9]+", "_", s).strip("_")


def _carpeta_import():
    ruta = os.path.join(current_app.root_path, "..", "uploads", "import")
    os.makedirs(ruta, exist_ok=True)
    return ruta


def _leer_archivo(file):
    """Devuelve (headers, filas) de un CSV o XLSX. Las filas son listas de texto."""
    nombre = (file.filename or "").lower()
    if nombre.endswith(".xlsx") or nombre.endswith(".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(file.read()), read_only=True, data_only=True)
        ws = wb.active
        datos = list(ws.iter_rows(values_only=True))
        if not datos:
            return [], []
        headers = [("" if h is None else str(h)) for h in datos[0]]
        filas = [[("" if v is None else str(v)) for v in r] for r in datos[1:]]
        return headers, filas

    raw = file.read()
    texto = None
    for enc in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            texto = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if texto is None:
        texto = raw.decode("latin-1", "replace")
    lineas = texto.splitlines()
    if not lineas:
        return [], []
    delim = ";" if lineas[0].count(";") >= lineas[0].count(",") else ","
    reader = csv.reader(io.StringIO(texto), delimiter=delim)
    filas_todas = list(reader)
    if not filas_todas:
        return [], []
    return filas_todas[0], filas_todas[1:]


def _sugerir_mapeo(headers, tipo):
    """Sugiere qué columna del archivo corresponde a cada campo destino."""
    headers_norm = [_norm(h) for h in headers]
    mapeo = {}
    for campo, _etq, _obl, claves in CAMPOS[tipo]:
        elegido = ""
        for i, hn in enumerate(headers_norm):
            if not hn:
                continue
            if any(hn == c or c in hn for c in claves):
                elegido = str(i)
                break
        mapeo[campo] = elegido
    return mapeo


def _num(valor, defecto=0.0):
    if valor is None:
        return defecto
    s = str(valor).strip().replace("$", "").replace(" ", "")
    # quita separador de miles y normaliza decimal
    if "," in s and "." in s:
        s = s.replace(",", "")
    elif s.count(",") == 1 and "." not in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return defecto


def _ent(valor, defecto=0):
    try:
        return int(round(_num(valor, defecto)))
    except (TypeError, ValueError):
        return defecto


@importar_bp.route("/", methods=["GET"])
@login_required
@role_required("administrador")
def importar_home():
    return render_template("importar.html", stage="upload")


@importar_bp.route("/analizar", methods=["POST"])
@login_required
@role_required("administrador")
def importar_analizar():
    tipo = request.form.get("tipo", "clientes")
    if tipo not in CAMPOS:
        flash("Tipo de importación no válido.", "danger")
        return redirect(url_for("importar.importar_home"))
    archivo = request.files.get("archivo")
    if not archivo or not archivo.filename:
        flash("Selecciona un archivo CSV o Excel.", "danger")
        return redirect(url_for("importar.importar_home"))

    try:
        headers, filas = _leer_archivo(archivo)
    except Exception as e:
        flash(f"No se pudo leer el archivo: {e}", "danger")
        return redirect(url_for("importar.importar_home"))

    if not headers or not filas:
        flash("El archivo no tiene datos (o no se detectaron columnas).", "danger")
        return redirect(url_for("importar.importar_home"))

    token = uuid.uuid4().hex
    with open(os.path.join(_carpeta_import(), token + ".json"), "w", encoding="utf-8") as f:
        json.dump({"headers": headers, "filas": filas, "tipo": tipo}, f)

    mapeo = _sugerir_mapeo(headers, tipo)
    preview = filas[:8]
    return render_template(
        "importar.html", stage="map", tipo=tipo, token=token,
        headers=headers, preview=preview, mapeo=mapeo,
        campos=CAMPOS[tipo], total=len(filas),
    )


def _valor(fila, indice):
    if indice is None or indice == "":
        return ""
    try:
        i = int(indice)
    except ValueError:
        return ""
    return fila[i].strip() if 0 <= i < len(fila) else ""


@importar_bp.route("/ejecutar", methods=["POST"])
@login_required
@role_required("administrador")
def importar_ejecutar():
    token = request.form.get("token", "")
    ruta = os.path.join(_carpeta_import(), token + ".json")
    if not os.path.exists(ruta):
        flash("La sesión de importación expiró. Vuelve a subir el archivo.", "danger")
        return redirect(url_for("importar.importar_home"))

    with open(ruta, encoding="utf-8") as f:
        datos = json.load(f)
    tipo = datos["tipo"]
    filas = datos["filas"]

    # Mapeo: campo -> índice de columna del archivo
    mapeo = {campo: request.form.get("map_" + campo, "") for campo, *_ in CAMPOS[tipo]}

    creados = actualizados = omitidos = 0

    if tipo == "clientes":
        for fila in filas:
            nombre = _valor(fila, mapeo.get("full_name"))
            if not nombre:
                omitidos += 1
                continue
            cedula = _valor(fila, mapeo.get("cedula")) or None
            existente = Customer.query.filter_by(cedula=cedula).first() if cedula else None
            if existente:
                if not existente.business_name:
                    existente.business_name = _valor(fila, mapeo.get("business_name")) or None
                if not existente.phone:
                    existente.phone = _valor(fila, mapeo.get("phone")) or None
                if not existente.email:
                    existente.email = _valor(fila, mapeo.get("email")) or None
                if not existente.address:
                    existente.address = _valor(fila, mapeo.get("address")) or None
                actualizados += 1
            else:
                db.session.add(Customer(
                    full_name=nombre[:150],
                    cedula=cedula,
                    business_name=_valor(fila, mapeo.get("business_name"))[:150] or None,
                    phone=_valor(fila, mapeo.get("phone"))[:20] or None,
                    email=_valor(fila, mapeo.get("email"))[:100] or None,
                    address=_valor(fila, mapeo.get("address"))[:200] or None,
                    customer_type="efectivo",
                    active=True,
                ))
                creados += 1
    else:  # productos
        for fila in filas:
            nombre = _valor(fila, mapeo.get("name"))
            if not nombre:
                omitidos += 1
                continue
            sku = _valor(fila, mapeo.get("sku")) or None
            existente = None
            if sku:
                existente = Product.query.filter_by(sku=sku).first()
            if existente:
                if _valor(fila, mapeo.get("cost_price")):
                    existente.cost_price = _num(_valor(fila, mapeo.get("cost_price")), float(existente.cost_price or 0))
                if _valor(fila, mapeo.get("sale_price")):
                    existente.sale_price = _num(_valor(fila, mapeo.get("sale_price")), float(existente.sale_price or 0))
                actualizados += 1
            else:
                p = Product(
                    name=nombre[:150],
                    sku=sku,
                    cost_price=_num(_valor(fila, mapeo.get("cost_price")), 0),
                    sale_price=_num(_valor(fila, mapeo.get("sale_price")), 0),
                    current_stock=_ent(_valor(fila, mapeo.get("current_stock")), 0),
                    unit_measure="unidad",
                    active=True,
                )
                db.session.add(p)
                creados += 1

    db.session.commit()
    try:
        os.remove(ruta)
    except OSError:
        pass

    registrar("importacion_datos",
              detalle=f"Importación {tipo}: {creados} nuevo(s), {actualizados} actualizado(s), {omitidos} omitido(s)")
    flash(f"Importación completada: {creados} nuevo(s), {actualizados} actualizado(s), {omitidos} omitido(s).", "success")
    if tipo == "clientes":
        return redirect(url_for("customers.customers_list"))
    return redirect(url_for("products.products_list"))
